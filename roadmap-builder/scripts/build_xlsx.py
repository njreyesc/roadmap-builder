#!/usr/bin/env python3
"""Рендерер roadmap в Excel.

Использование: python build_xlsx.py roadmap.json roadmap.xlsx
"""
import sys

from openpyxl import Workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from roadmap_common import STYLE, flag_markers, flatten_rows, group_label, load_roadmap

COL_GROUP = 1   # A — команда/направление
COL_LABEL = 2   # B — подпись строки
FIRST_WEEK_COL = 3
HEADER_ROWS = 2  # месяц, неделя
FIRST_DATA_ROW = HEADER_ROWS + 2  # +1 строка под title


def fill(hexcode):
    return PatternFill("solid", fgColor=hexcode)


def thin_border(color=STYLE["grid_line"]):
    side = Side(style="thin", color=color)
    return Border(left=side, right=side, top=side, bottom=side)


def merge_range_at(ws, r, c):
    """Merge-диапазон, накрывающий (r, c), либо None."""
    for mr in ws.merged_cells.ranges:
        if mr.min_row <= r <= mr.max_row and mr.min_col <= c <= mr.max_col:
            return mr
    return None


def anchor_cell(ws, r, c):
    """Записываемая ячейка для (r, c): якорь merge-диапазона либо сама ячейка.

    В MergedCell нельзя писать value/comment — только в верхнюю-левую ячейку.
    """
    cell = ws.cell(row=r, column=c)
    if isinstance(cell, MergedCell):
        mr = merge_range_at(ws, r, c)
        if mr:
            return ws.cell(row=mr.min_row, column=mr.min_col)
    return cell


def set_col_side(ws, col, which, color, first_row, last_row):
    """Ставит вертикальную границу (which='left'|'right') колонки col по строкам.

    openpyxl при сохранении переписывает границы merge-диапазона от якорной
    ячейки, поэтому для объединённых полос границу ставим на якорь — и только
    если col совпадает с нужным краем merge-диапазона (внутри merge Excel
    линию не отрисует). Используется для линии «сегодня» и флажков границ.
    """
    side = Side(style="medium", color=color)
    for rr in range(first_row, last_row):
        mr = merge_range_at(ws, rr, col)
        if mr is not None:
            if which == "right" and mr.max_col != col:
                continue
            if which == "left" and mr.min_col != col:
                continue
            cell = ws.cell(row=mr.min_row, column=mr.min_col)
        else:
            cell = ws.cell(row=rr, column=col)
        b = cell.border
        sides = dict(left=b.left, right=b.right, top=b.top, bottom=b.bottom)
        sides[which] = side
        cell.border = Border(**sides)


def main(json_path, out_path):
    data, tl = load_roadmap(json_path)
    wb = Workbook()
    ws = wb.active
    ws.title = "Roadmap"

    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)

    # Заголовок
    ws.cell(row=1, column=1, value=data["title"]).font = Font(bold=True, size=14)
    if data.get("subtitle"):
        c = ws.cell(row=1, column=FIRST_WEEK_COL, value=data["subtitle"])
        c.font = Font(italic=True, size=10, color="4472C4")

    # Шапка: месяцы (объединённые) и недели
    today_col = None
    if data.get("today"):
        ti = tl.week_index(data["today"], clamp=False)
        if ti is not None:
            today_col = FIRST_WEEK_COL + ti

    for label, i0, i1 in tl.month_spans():
        c0, c1 = FIRST_WEEK_COL + i0, FIRST_WEEK_COL + i1
        ws.merge_cells(start_row=2, start_column=c0, end_row=2, end_column=c1)
        cell = ws.cell(row=2, column=c0, value=label)
        cell.fill = fill(STYLE["header_fill"])
        cell.font = Font(bold=True, color=STYLE["header_font"])
        cell.alignment = center
    for i in range(tl.n):
        col = FIRST_WEEK_COL + i
        cell = ws.cell(row=3, column=col, value=tl.week_label(i))
        cell.fill = fill(STYLE["today_line"] if col == today_col else STYLE["week_fill"])
        if col == today_col:
            cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = center
        cell.border = thin_border()
        ws.column_dimensions[get_column_letter(col)].width = 8.5

    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 30

    rows = flatten_rows(data)
    # Сетка + данные
    bar_styles = {
        "work": (STYLE["bar_work_fill"], STYLE["bar_work_line"]),
        "estimate": (STYLE["bar_estimate_fill"], STYLE["bar_estimate_line"]),
        "analytics": (STYLE["bar_analytics_fill"], STYLE["bar_analytics_line"]),
        "testing": (STYLE["bar_testing_fill"], STYLE["bar_testing_line"]),
        "vacation": (STYLE["bar_vacation_fill"], STYLE["bar_vacation_line"]),
    }
    r = FIRST_DATA_ROW
    group_start = {}
    for g, row, is_first in rows:
        if is_first:
            group_start[id(g)] = r
        # фон сетки
        for i in range(tl.n):
            cell = ws.cell(row=r, column=FIRST_WEEK_COL + i)
            cell.border = thin_border()
        lab = ws.cell(row=r, column=COL_LABEL, value=row.get("label", ""))
        lab.alignment = left
        lab.font = Font(size=9)
        lab.border = thin_border()
        ws.row_dimensions[r].height = 22

        occupied = set()
        for item in row.get("items", []):
            t = item.get("type")
            if t == "bar":
                clip = tl.clip_bar(item["start"], item["end"], item.get("label", ""))
                if clip is None:
                    print(f"! полоса '{item.get('label')}' в строке '{row.get('label')}' "
                          f"целиком вне диапазона roadmap — пропущена")
                    continue
                i0, i1, cut_l, cut_r = clip
                label = item.get("label", "")
                if cut_r:
                    label = (label + " →") if label else "→"
                if cut_l:
                    label = ("← " + label) if label else "←"
                c0, c1 = FIRST_WEEK_COL + i0, FIRST_WEEK_COL + i1
                if any(c in occupied for c in range(c0, c1 + 1)):
                    print(f"! перекрытие в строке '{row.get('label')}', полоса '{item.get('label')}' обрезана до свободных недель")
                    while c0 <= c1 and c0 in occupied:  # сдвигаем начало до свободной колонки
                        c0 += 1
                    end = c0
                    while end + 1 <= c1 and end + 1 not in occupied:
                        end += 1
                    c1 = end
                    if c0 > c1:
                        print(f"!   свободных недель не осталось — полоса пропущена")
                        continue
                if c1 > c0:
                    ws.merge_cells(start_row=r, start_column=c0, end_row=r, end_column=c1)
                bar_fill, bar_line = bar_styles.get(item.get("style", "work"), bar_styles["work"])
                cell = ws.cell(row=r, column=c0, value=label)
                cell.alignment = center
                cell.font = Font(size=9, color=STYLE["text"])
                for c in range(c0, c1 + 1):
                    ws.cell(row=r, column=c).fill = fill(bar_fill)
                    ws.cell(row=r, column=c).border = thin_border(bar_line)
                    occupied.add(c)
            elif t == "milestone":
                i0 = tl.week_index(item["date"], clamp=False)
                if i0 is None:
                    print(f"! веха '{item.get('label', item['date'])}' в строке "
                          f"'{row.get('label')}' вне диапазона roadmap — пропущена")
                    continue
                c0 = FIRST_WEEK_COL + i0
                status = item.get("status", "done")
                color = STYLE["milestone_done_fill"] if status == "done" else STYLE["milestone_planned_line"]
                text = "◆" if status == "done" else "◇"
                if item.get("label"):
                    text += " " + item["label"]
                if c0 in occupied or isinstance(ws.cell(row=r, column=c0), MergedCell):
                    # веха попала на полосу — дописываем её к подписи полосы
                    print(f"! веха '{item.get('label', item['date'])}' в строке "
                          f"'{row.get('label')}' перекрывает полосу — добавлена к её подписи")
                    a = anchor_cell(ws, r, c0)
                    a.value = f"{a.value} {text}" if a.value else text
                else:
                    cell = ws.cell(row=r, column=c0, value=text)
                    cell.font = Font(size=12, color=color, bold=True)
                    cell.alignment = center
                occupied.add(c0)
            elif t == "note":
                i0 = tl.week_index(item["date"], clamp=False)
                if i0 is None:
                    print(f"! выноска '{item['text'][:30]}…' вне диапазона roadmap — пропущена")
                    continue
                c0 = FIRST_WEEK_COL + i0
                cell = ws.cell(row=r, column=c0)
                # выноска — комментарий + голубая метка, чтобы не воевать за место с полосами
                if not isinstance(cell, MergedCell) and cell.value is None and c0 not in occupied:
                    cell.value = "💬"
                    cell.alignment = center
                    cell.fill = fill(STYLE["note_fill"])
                # комментарий можно повесить только на якорную ячейку merge-диапазона
                anchor_cell(ws, r, c0).comment = Comment(item["text"], "roadmap", height=80, width=260)
        r += 1

    # Объединение колонки групп
    for g in data["groups"]:
        start = group_start[id(g)]
        n = max(1, len(g.get("rows") or [1]))
        if n > 1:
            ws.merge_cells(start_row=start, start_column=COL_GROUP, end_row=start + n - 1, end_column=COL_GROUP)
        cell = ws.cell(row=start, column=COL_GROUP, value=group_label(g))
        cell.font = Font(bold=True, size=10)
        cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        for rr in range(start, start + n):
            ws.cell(row=rr, column=COL_GROUP).fill = fill(STYLE["group_fill"])
            ws.cell(row=rr, column=COL_GROUP).border = thin_border()

    # Линия "сегодня": правая граница колонки текущей недели.
    if today_col:
        set_col_side(ws, today_col, "right", STYLE["today_line"], 3, r)

    # Флажки границ roadmap: цветной «шест» (левая граница первой недели —
    # старт, правая граница последней — финиш) через все строки + подпись с
    # датой в шапке недели.
    start_col = FIRST_WEEK_COL
    end_col = FIRST_WEEK_COL + tl.n - 1
    flags = flag_markers(data)
    for kind, flabel in flags:
        col = start_col if kind == "start" else end_col
        which = "left" if kind == "start" else "right"
        cell = ws.cell(row=3, column=col, value=f"⚑ {flabel}\n{tl.week_label(col - FIRST_WEEK_COL)}")
        cell.fill = fill(STYLE[f"flag_{kind}_fill"])
        cell.font = Font(bold=True, color="FFFFFF", size=9)
        cell.alignment = center
        cell.border = thin_border()
        # «шест» ставим после подписи — цветная граница выигрывает и на шапке
        set_col_side(ws, col, which, STYLE[f"flag_{kind}_line"], 3, r)
    if flags:
        ws.row_dimensions[3].height = 30  # под подпись флажка в две строки

    # Легенда
    lr = r + 2
    ws.cell(row=lr, column=COL_LABEL, value="Легенда:").font = Font(bold=True, size=9)
    legend = [
        ("◆ выполнено / факт", STYLE["milestone_done_fill"], None),
        ("◇ план / под вопросом", STYLE["milestone_planned_line"], None),
        ("работы (разработка)", None, STYLE["bar_work_fill"]),
        ("аналитика (SDD)", None, STYLE["bar_analytics_fill"]),
        ("тестирование", None, STYLE["bar_testing_fill"]),
        ("оценка т/з", None, STYLE["bar_estimate_fill"]),
        ("отпуск", None, STYLE["bar_vacation_fill"]),
        ("⚑ старт", "FFFFFF", STYLE["flag_start_fill"]),
        ("⚑ финиш", "FFFFFF", STYLE["flag_end_fill"]),
    ]
    for k, (text, font_color, bg) in enumerate(legend):
        cell = ws.cell(row=lr + 1 + k, column=COL_LABEL, value=text)
        cell.font = Font(size=9, color=font_color or STYLE["text"])
        if bg:
            cell.fill = fill(bg)

    ws.freeze_panes = ws.cell(row=FIRST_DATA_ROW, column=FIRST_WEEK_COL)
    wb.save(out_path)
    print(f"OK: {out_path} — {len(rows)} строк, {tl.n} недель")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit("Использование: python build_xlsx.py roadmap.json roadmap.xlsx")
    main(sys.argv[1], sys.argv[2])
