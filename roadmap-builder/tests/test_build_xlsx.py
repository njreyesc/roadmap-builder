"""Тесты рендерера Excel (build_xlsx): флажки границ roadmap."""
import json

from openpyxl import load_workbook

import build_xlsx
from build_xlsx import FIRST_WEEK_COL
from roadmap_common import STYLE


def render(tmp_path, data):
    src = tmp_path / "roadmap.json"
    out = tmp_path / "roadmap.xlsx"
    src.write_text(json.dumps(data), encoding="utf-8")
    build_xlsx.main(str(src), str(out))
    return load_workbook(str(out))["Roadmap"]


def one_week_roadmap():
    return {
        "title": "К7",
        "start": "2026-01-05",
        "end": "2026-01-11",
        "groups": [{"name": "Ф", "rows": [{"label": "Р", "items": [
            {"type": "bar", "start": "2026-01-05", "end": "2026-01-11", "label": "1 нед"},
        ]}]}],
    }


def test_one_week_flags_share_cell_with_merged_label(tmp_path):
    # roadmap в одну неделю: старт и финиш в одной колонке — подписи
    # объединяются, обе medium-границы («шесты») сохраняются
    ws = render(tmp_path, one_week_roadmap())
    cell = ws.cell(row=3, column=FIRST_WEEK_COL)
    assert "Старт" in cell.value and "Финиш" in cell.value
    assert cell.value.startswith("⚑ ")
    assert cell.border.left.style == "medium"
    assert cell.border.left.color.rgb.endswith(STYLE["flag_start_line"])
    assert cell.border.right.style == "medium"
    assert cell.border.right.color.rgb.endswith(STYLE["flag_end_line"])


def test_one_week_flag_poles_reach_data_rows(tmp_path):
    # «шесты» обоих флажков идут и через строки данных, не только шапку
    ws = render(tmp_path, one_week_roadmap())
    cell = ws.cell(row=build_xlsx.FIRST_DATA_ROW, column=FIRST_WEEK_COL)
    assert cell.border.left.style == "medium"
    assert cell.border.right.style == "medium"


def test_multi_week_flags_stay_separate(tmp_path):
    data = one_week_roadmap()
    data["end"] = "2026-01-18"  # 2 недели
    ws = render(tmp_path, data)
    start = ws.cell(row=3, column=FIRST_WEEK_COL)
    end = ws.cell(row=3, column=FIRST_WEEK_COL + 1)
    assert "Старт" in start.value and "Финиш" not in start.value
    assert "Финиш" in end.value and "Старт" not in end.value
    assert start.border.left.style == "medium"
    assert start.border.left.color.rgb.endswith(STYLE["flag_start_line"])
    assert end.border.right.style == "medium"
    assert end.border.right.color.rgb.endswith(STYLE["flag_end_line"])


def test_one_week_custom_labels_merged(tmp_path):
    data = one_week_roadmap()
    data["start_label"] = "Kickoff"
    data["end_label"] = "Release"
    ws = render(tmp_path, data)
    assert ws.cell(row=3, column=FIRST_WEEK_COL).value.startswith("⚑ Kickoff / Release")
